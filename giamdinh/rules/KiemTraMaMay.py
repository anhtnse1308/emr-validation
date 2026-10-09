"""
Rule – Kiểm tra MA_MAY trong XML3 đối với các DVKT yêu cầu khai báo máy
Căn cứ: Danh mục DVKT yêu cầu MA_MAY (file Excel ngoài, ~4800 dòng)

Logic:
  - Load file Excel danh mục DVKT cần MA_MAY (cột MA_DVKT)
  - Với mỗi dòng XML3: nếu MA_DICH_VU khớp MA_DVKT trong danh mục
    → kiểm tra MA_MAY có giá trị (sau khi trim khoảng trắng) không
  - Nếu MA_MAY trống → cảnh báo WARN.MM01

File danh mục Excel cần có cột: MA_DVKT (bắt buộc)
Các cột TEN_DVKT_TT23, TEN_DVKT_PHE_DUYET dùng để hiển thị trong báo cáo.

Dùng:
    loader = DanhMucMaMayLoader("path/to/DanhMucDVKT_MaMay.xlsx")
    rule   = KiemTraMaMay(loader)
    errors = rule.check(all_objects)
    out    = KiemTraMaMay.write_excel(excel_file, errors)
"""

from __future__ import annotations
import os
from datetime import datetime
from giamdinh.GiamDinhBase import GiamDinhBase, GiamDinhError

CAN_CU = "Danh mục DVKT yêu cầu khai báo mã máy thực hiện"


# ---------------------------------------------------------------------------
# Loader danh mục
# ---------------------------------------------------------------------------
class DanhMucMaMayLoader:
    """
    Load file Excel danh mục DVKT yêu cầu MA_MAY.
    Cột bắt buộc: MA_DVKT
    Cột tùy chọn: TEN_DVKT_TT23, TEN_DVKT_PHE_DUYET
    """

    def __init__(self, filepath: str):
        if not os.path.exists(filepath):
            raise FileNotFoundError(f"Không tìm thấy file danh mục: {filepath}")
        self.filepath = filepath
        # MA_DVKT → {ten_tt23, ten_phe_duyet}
        self._catalog: dict[str, dict] = {}
        self._load()

    def _load(self):
        try:
            import pandas as pd
        except ImportError:
            raise ImportError("Cần cài pandas: pip install pandas openpyxl")

        df = pd.read_excel(self.filepath, dtype=str).fillna("")

        # Tìm cột MA_DVKT (không phân biệt hoa thường)
        col_map = {c.strip().upper(): c for c in df.columns}
        col_ma = col_map.get("MA_DVKT")
        if not col_ma:
            raise ValueError(f"File danh mục thiếu cột MA_DVKT: {self.filepath}")

        col_ten_tt23    = col_map.get("TEN_DVKT_TT23", "")
        col_ten_pd      = col_map.get("TEN_DVKT_PHE_DUYET", "")

        for _, row in df.iterrows():
            ma = str(row[col_ma]).strip()
            if not ma or ma.lower() in ("nan", "none", ""):
                continue
            self._catalog[ma] = {
                "ten_tt23":  str(row[col_ten_tt23]).strip() if col_ten_tt23 else "",
                "ten_pd":    str(row[col_ten_pd]).strip()   if col_ten_pd   else "",
            }

    def requires_ma_may(self, ma_dvkt: str) -> bool:
        return ma_dvkt.strip() in self._catalog

    def get_info(self, ma_dvkt: str) -> dict:
        return self._catalog.get(ma_dvkt.strip(), {})

    @property
    def total(self) -> int:
        return len(self._catalog)

    def __repr__(self) -> str:
        return (f"DanhMucMaMayLoader("
                f"file={os.path.basename(self.filepath)!r}, "
                f"dvkt={self.total})")


# ---------------------------------------------------------------------------
# Rule
# ---------------------------------------------------------------------------
class KiemTraMaMay(GiamDinhBase):
    """
    Kiểm tra MA_MAY trong XML3 với các DVKT thuộc danh mục yêu cầu khai báo máy.
    WARN.MM01 – MA_DICH_VU yêu cầu MA_MAY nhưng MA_MAY trống hoặc chỉ khoảng trắng.
    """

    def __init__(self, loader: DanhMucMaMayLoader):
        self.loader = loader

    def check(self, all_objects: dict) -> list[GiamDinhError]:
        errors: list[GiamDinhError] = []

        for row_excel, rec in self._get_rows(all_objects, "XML3"):
            ma_lk    = (getattr(rec, "MA_LK",      None) or "").strip()
            ma_dv    = (getattr(rec, "MA_DICH_VU",  None) or "").strip()
            ma_may   = (getattr(rec, "MA_MAY",      None) or "")
            ten_dv   = (getattr(rec, "TEN_DICH_VU", None) or "").strip()

            if not ma_lk or not ma_dv:
                continue
            if not self.loader.requires_ma_may(ma_dv):
                continue

            # Trim khoảng trắng, kiểm tra có giá trị không
            if not str(ma_may).strip():
                info = self.loader.get_info(ma_dv)
                ten  = info.get("ten_tt23") or info.get("ten_pd") or ten_dv or ma_dv
                errors.append(GiamDinhError(
                    sheet="XML3",
                    ma_lk=ma_lk,
                    row_excel=row_excel,
                    ma_ly_do="WARN.MM01",
                    mo_ta=(
                        f"DVKT '{ten}' (mã: {ma_dv}) yêu cầu khai báo MA_MAY "
                        f"nhưng MA_MAY để trống."
                    ),
                    can_cu=CAN_CU,
                    ma_dich_vu=ma_dv,
                ))

        return errors

    # -----------------------------------------------------------------------
    # Export Excel riêng
    # -----------------------------------------------------------------------
    @staticmethod
    def write_excel(excel_file: str, errors: list[GiamDinhError]) -> str:
        """
        Xuất kết quả kiểm tra MA_MAY.
        Tên: <ten_goc>[<timestamp>]_mamay.xlsx

        Sheet TONG_HOP – tổng số lỗi
        Sheet CHI_TIET – toàn bộ lỗi: STT, ROW_EXCEL, MA_LK, MA_DVKT, MO_TA
        """
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        FILL_HDR  = PatternFill("solid", start_color="1F3864")
        FILL_ERR  = PatternFill("solid", start_color="FCE4D6")
        FILL_EVEN = PatternFill("solid", start_color="FFFFFF")
        FILL_OK   = PatternFill("solid", start_color="E2EFDA")
        FILL_SUM  = PatternFill("solid", start_color="C00000")

        FONT_HDR   = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        FONT_DATA  = Font(name="Arial", size=10)
        FONT_RED   = Font(name="Arial", size=10, color="C00000", bold=True)
        FONT_WHITE = Font(name="Arial", size=10, color="FFFFFF", bold=True)

        ALIGN_CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ALIGN_L   = Alignment(vertical="center", wrap_text=True)
        ALIGN_TOP = Alignment(vertical="top", wrap_text=True)

        _THIN = Side(style="thin", color="AAAAAA")
        BDR   = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

        def hdr(ws, r, c, v):
            cell = ws.cell(row=r, column=c, value=v)
            cell.fill = FILL_HDR; cell.font = FONT_HDR
            cell.alignment = ALIGN_CTR; cell.border = BDR
            return cell

        def dat(ws, r, c, v, fill=FILL_EVEN, font=FONT_DATA, align=ALIGN_L):
            cell = ws.cell(row=r, column=c, value=v)
            cell.fill = fill; cell.font = font
            cell.alignment = align; cell.border = BDR
            return cell

        def auto_w(ws, headers, start=2):
            for ci, h in enumerate(headers, 1):
                mx = len(str(h))
                for ri in range(start, ws.max_row + 1):
                    mx = max(mx, min(len(str(ws.cell(ri, ci).value or "").split("\n")[0]), 80))
                ws.column_dimensions[get_column_letter(ci)].width = mx + 3

        wb = Workbook()
        if wb.active:
            wb.remove(wb.active)

        now_str = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

        # TONG_HOP
        ws = wb.create_sheet("TONG_HOP")
        ws.merge_cells("A1:D1")
        tc = ws.cell(1, 1, f"KIỂM TRA MA_MAY DVKT  –  {now_str}")
        tc.fill = FILL_HDR
        tc.font = Font(bold=True, color="FFFFFF", name="Arial", size=12)
        tc.alignment = ALIGN_CTR; ws.row_dimensions[1].height = 28

        for ci, h in enumerate(["MÃ CẢNH BÁO", "Ý NGHĨA", "SỐ DÒNG"], 1):
            hdr(ws, 2, ci, h)

        so = len(errors)
        fill = FILL_ERR if so else FILL_OK
        dat(ws, 3, 1, "WARN.MM01", fill=fill, align=ALIGN_CTR)
        dat(ws, 3, 2, "DVKT yêu cầu MA_MAY nhưng để trống", fill=fill)
        dat(ws, 3, 3, so, fill=fill,
            font=FONT_RED if so else FONT_DATA, align=ALIGN_CTR)

        dat(ws, 4, 1, "TỔNG",  fill=FILL_SUM, font=FONT_WHITE, align=ALIGN_CTR)
        dat(ws, 4, 2, "",      fill=FILL_SUM)
        dat(ws, 4, 3, so,      fill=FILL_SUM, font=FONT_WHITE, align=ALIGN_CTR)

        if not errors:
            ws.merge_cells("A6:D6")
            ok = ws.cell(6, 1, "✓ Tất cả DVKT yêu cầu MA_MAY đều đã khai báo.")
            ok.fill = FILL_OK
            ok.font = Font(name="Arial", size=11, bold=True, color="375623")
            ok.alignment = ALIGN_CTR

        auto_w(ws, ["MÃ CẢNH BÁO", "Ý NGHĨA", "SỐ DÒNG"], start=3)
        ws.freeze_panes = "A3"

        # CHI_TIET
        ws_ct = wb.create_sheet("CHI_TIET")
        CT_HDRS = ["STT", "ROW_EXCEL", "MA_LK", "MA_DVKT", "MÔ TẢ"]
        for ci, h in enumerate(CT_HDRS, 1):
            hdr(ws_ct, 1, ci, h)

        for idx, err in enumerate(errors, 1):
            ri   = idx + 1
            fill = FILL_ERR if idx % 2 else FILL_EVEN
            dat(ws_ct, ri, 1, idx,           fill=fill, align=ALIGN_CTR)
            dat(ws_ct, ri, 2, err.row_excel,  fill=fill, align=ALIGN_CTR)
            dat(ws_ct, ri, 3, err.ma_lk,     fill=fill)
            dat(ws_ct, ri, 4, err.ma_dich_vu, fill=fill)
            dat(ws_ct, ri, 5, err.mo_ta,     fill=fill, align=ALIGN_TOP)
            ws_ct.row_dimensions[ri].height = max(25, len(err.mo_ta) // 60 * 14 + 14)

        auto_w(ws_ct, CT_HDRS)
        ws_ct.freeze_panes = "A2"
        ws_ct.row_dimensions[1].height = 28

        # Lưu
        folder = os.path.dirname(os.path.abspath(excel_file))
        base   = os.path.splitext(os.path.basename(excel_file))[0]
        ts     = datetime.now().strftime("%Y%m%d_%H%M%S")
        out    = os.path.join(folder, f"{base}[{ts}]_mamay.xlsx")
        wb.save(out)
        return out
